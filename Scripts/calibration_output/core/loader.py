import datetime as dt

import pandas as pd
from openpyxl import load_workbook

import config

LOG_COLUMNS = {
    1: "category",
    2: "key",
    3: "date",
    4: "camera",
    5: "frame_count",
    6: "world_error",
    7: "image_error",
}

META_COLUMNS = {
    config.DAILY_META_SHEET: {
        1: "category",
        2: "key",
        3: "date",
        4: "temperature",
        5: "start_time",
        6: "technique",
        7: "auto_stop",
        22: "notes_a",
    },
    config.ERROR_META_SHEET: {
        1: "category",
        2: "key",
        3: "date",
        4: "temperature",
        5: "start_time",
        6: "technique",
        7: "auto_stop",
        16: "notes_a",
        17: "notes_b",
    },
    config.TRIAL_META_SHEET: {
        1: "category",
        2: "key",
        3: "date",
        4: "temperature",
        5: "start_time",
        6: "technique",
        7: "auto_stop",
        26: "cameras_relative",
        27: "wand_wave_time",
        28: "notes_a",
        29: "notes_b",
        30: "notes_c",
    },
}

META_FIELDS = [
    "temperature",
    "start_time",
    "technique",
    "auto_stop",
    "cameras_relative",
    "wand_wave_time",
    "notes_a",
    "notes_b",
    "notes_c",
]

CACHE_COLUMNS = [
    "category",
    "key",
    "date",
    "camera",
    "camera_id",
    "camera_index",
    "frame_count",
    "world_error",
    "image_error",
] + META_FIELDS


EXCEL_ERRORS = {
    "#REF!",
    "#VALUE!",
    "#N/A",
    "#DIV/0!",
    "#NAME?",
    "#NULL!",
    "#NUM!",
    "#SPILL!",
    "#CALC!",
}


def _clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        if not text or text.upper() in EXCEL_ERRORS:
            return None
        return text
    return value


def _as_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return None


def _as_time(value):
    if isinstance(value, dt.time):
        return value.strftime("%H:%M")
    if isinstance(value, dt.datetime):
        return value.strftime("%H:%M")
    return None


def _as_number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _read_log(workbook):
    sheet = workbook[config.LOG_SHEET]
    rows = []
    for row in sheet.iter_rows(min_row=config.LOG_HEADER_ROW + 1, max_col=max(LOG_COLUMNS)):
        record = {}
        for index, name in LOG_COLUMNS.items():
            record[name] = _clean(row[index - 1].value)
        if record["date"] is None or record["camera"] is None:
            continue
        record["date"] = _as_date(record["date"])
        if record["date"] is None:
            continue
        record["frame_count"] = _as_number(record["frame_count"])
        record["world_error"] = _as_number(record["world_error"])
        record["image_error"] = _as_number(record["image_error"])
        rows.append(record)
    return pd.DataFrame(rows)


def _read_meta(workbook):
    rows = []
    for sheet_name, mapping in META_COLUMNS.items():
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(min_row=2, max_col=max(mapping)):
            record = {}
            for index, name in mapping.items():
                record[name] = _clean(row[index - 1].value)
            if record.get("date") is None or record.get("category") is None:
                continue
            record["date"] = _as_date(record["date"])
            if record["date"] is None:
                continue
            record["start_time"] = _as_time(record.get("start_time"))
            record["wand_wave_time"] = _as_time(record.get("wand_wave_time")) or record.get(
                "wand_wave_time"
            )
            record["temperature"] = _as_number(record.get("temperature"))
            record["auto_stop"] = _as_number(record.get("auto_stop"))
            for field in META_FIELDS:
                record.setdefault(field, None)
            rows.append(record)
    frame = pd.DataFrame(rows)
    if frame.empty:
        return frame
    frame = frame.drop_duplicates(subset=["category", "key", "date"], keep="first")
    return frame


def _read_cameras(workbook):
    sheet = workbook[config.CAMERA_LIST_SHEET]
    column, start, end = config.CAMERA_LIST_RANGE
    names = []
    for row in range(start, end + 1):
        value = _clean(sheet[f"{column}{row}"].value)
        if value:
            names.append(value)
    return names


def _camera_id(name):
    return str(name).split("(")[0].strip()


def _camera_index(name):
    token = _camera_id(name).lstrip("#")
    try:
        return int(token)
    except ValueError:
        return 0


def build_cache(source_path=None, cache_path=None):
    source = source_path or config.TRACKER_PATH
    target = cache_path or config.CACHE_PATH
    suffix = str(source).lower()
    if suffix.endswith(".csv"):
        frame = pd.read_csv(source, keep_default_na=False, na_values=[""])
        for column in ("frame_count", "world_error", "image_error", "temperature", "auto_stop"):
            if column in frame.columns:
                frame[column] = pd.to_numeric(frame[column], errors="coerce")
        frame["date"] = pd.to_datetime(frame["date"], dayfirst=True, errors="coerce").dt.date
        cameras = sorted(frame["camera"].dropna().unique(), key=_camera_index)
    else:
        workbook = load_workbook(source, data_only=True, read_only=True)
        log = _read_log(workbook)
        meta = _read_meta(workbook)
        cameras = _read_cameras(workbook)
        workbook.close()
        if meta.empty:
            frame = log
            for field in META_FIELDS:
                frame[field] = None
        else:
            frame = log.merge(meta, on=["category", "key", "date"], how="left")
    for field in META_FIELDS:
        if field not in frame.columns:
            frame[field] = None
    frame["camera_id"] = frame["camera"].map(_camera_id)
    frame["camera_index"] = frame["camera"].map(_camera_index)
    frame = frame[CACHE_COLUMNS]
    frame = frame.sort_values(["date", "category", "key", "camera_index"], kind="stable")
    frame.to_csv(target, index=False)
    return frame, cameras


def load(source_path=None, cache_path=None):
    frame, cameras = build_cache(source_path, cache_path)
    frame["date"] = pd.to_datetime(frame["date"])
    if not cameras:
        cameras = sorted(frame["camera"].dropna().unique(), key=_camera_index)
    return frame, cameras
